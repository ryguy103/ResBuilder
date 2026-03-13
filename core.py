"""
core.py - Core functions for resbuilder resume/cover letter generation
Shared between CLI (prepare.py) and web interface (web/app.py)
"""

import os
import re
import json
from datetime import datetime
from pathlib import Path
from typing import Optional, Generator

from dotenv import load_dotenv

load_dotenv(Path(__file__).parent / ".env")

try:
    import yaml
    HAS_YAML = True
except ImportError:
    HAS_YAML = False

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
except ImportError:
    Document = None

try:
    from anthropic import Anthropic
except ImportError:
    Anthropic = None

try:
    import requests
    from bs4 import BeautifulSoup
    HAS_SCRAPING = True
except ImportError:
    HAS_SCRAPING = False


# Base path for the project (where resume.md lives)
BASE_PATH = Path(__file__).parent


def _copy_example_files():
    """On first run, copy .example files so new users have a starting point."""
    import shutil
    for name in ("profile.yaml", "resume.md"):
        target = BASE_PATH / name
        source = BASE_PATH / f"{name}.example"
        if not target.exists() and source.exists():
            shutil.copy2(source, target)

_copy_example_files()


def get_api_key() -> str:
    """Get Anthropic API key from environment"""
    key = os.environ.get("ANTHROPIC_API_KEY")
    if not key:
        raise ValueError("ANTHROPIC_API_KEY environment variable not set")
    return key


def load_profile() -> dict:
    """Load the master professional profile from profile.yaml"""
    profile_path = BASE_PATH / "profile.yaml"
    
    if not profile_path.exists():
        return {}
    
    if not HAS_YAML:
        return {}
    
    with open(profile_path, "r") as f:
        return yaml.safe_load(f) or {}


def save_profile(profile: dict) -> None:
    """Save the professional profile to profile.yaml"""
    if not HAS_YAML:
        raise ImportError("PyYAML not installed")
    
    profile_path = BASE_PATH / "profile.yaml"
    with open(profile_path, "w") as f:
        yaml.dump(profile, f, default_flow_style=False, allow_unicode=True, sort_keys=False)


def get_profile_summary(profile: dict, summary_type: str = "default") -> str:
    """Get a specific summary variation from the profile"""
    summaries = profile.get("personal", {}).get("summaries", {})
    return summaries.get(summary_type, summaries.get("default", ""))


def format_profile_for_ai(profile: dict) -> str:
    """Format the profile into a comprehensive context string for AI"""
    if not profile:
        return read_master_resume()
    
    lines = []
    personal = profile.get("personal", {})
    
    # Header
    lines.append(f"CANDIDATE: {personal.get('name', 'Unknown')}")
    lines.append(f"Location: {personal.get('location', '')}")
    lines.append(f"Email: {personal.get('email', '')}")
    lines.append(f"LinkedIn: {personal.get('linkedin', '')}")
    lines.append("")
    
    # Summaries
    summaries = personal.get("summaries", {})
    if summaries:
        lines.append("PROFESSIONAL SUMMARIES (use most relevant):")
        for key, summary in summaries.items():
            lines.append(f"  [{key}]: {summary.strip()}")
        lines.append("")
    
    # Experience
    lines.append("WORK EXPERIENCE:")
    for job in profile.get("experience", []):
        lines.append(f"\n{job.get('company', '')} - {job.get('location', '')}")
        for role in job.get("roles", []):
            lines.append(f"  {role.get('title', '')} ({role.get('dates', '')})")
            for highlight in role.get("highlights", []):
                lines.append(f"    - {highlight}")
    lines.append("")
    
    # Skills
    lines.append("SKILLS:")
    for category in profile.get("skills", {}).get("categories", []):
        items = ", ".join(category.get("items", []))
        lines.append(f"  {category.get('name', '')}: {items}")
    lines.append("")
    
    # Projects
    lines.append("PROJECTS:")
    for project in profile.get("projects", []):
        lines.append(f"  {project.get('name', '')} ({project.get('type', '')})")
        lines.append(f"    {project.get('description', '').strip()}")
        for achievement in project.get("achievements", []):
            lines.append(f"    - {achievement}")
    lines.append("")
    
    # Awards
    lines.append("AWARDS:")
    for award in profile.get("awards", []):
        lines.append(f"  {award.get('name', '')} ({award.get('date', '')})")
        lines.append(f"    {award.get('description', '').strip()}")
    lines.append("")
    
    # Key Metrics
    lines.append("KEY METRICS:")
    for metric in profile.get("key_metrics", []):
        lines.append(f"  - {metric.get('metric', '')}: {metric.get('value', '')} ({metric.get('context', '')})")
    lines.append("")
    
    # Education
    lines.append("EDUCATION:")
    for edu in profile.get("education", []):
        lines.append(f"  {edu.get('degree', '')} - {edu.get('institution', '')} ({edu.get('years', '')})")
    
    return "\n".join(lines)


def read_master_resume() -> str:
    """Read the master resume content (falls back if no profile)"""
    # Try profile first
    profile = load_profile()
    if profile:
        return format_profile_for_ai(profile)
    
    # Fall back to resume.md
    resume_path = BASE_PATH / "resume.md"
    if not resume_path.exists():
        raise FileNotFoundError(
            "No profile.yaml or resume.md found. "
            "Set up your profile at http://localhost:8000/profile or copy the example: "
            "cp profile.yaml.example profile.yaml"
        )
    with open(resume_path, "r") as f:
        return f.read()


def get_company_slug(company: str) -> str:
    """Convert company name to folder-safe slug"""
    return company.lower().replace(" ", "-").replace(".", "")


def get_companies_dir() -> Path:
    """Get the companies directory path"""
    return BASE_PATH / "companies"


def list_applications() -> list:
    """List all past applications with their files"""
    companies_dir = get_companies_dir()
    applications = []
    
    if not companies_dir.exists():
        return applications
    
    for company_dir in sorted(companies_dir.iterdir(), reverse=True):
        if not company_dir.is_dir():
            continue
        
        # Get all files in the company folder
        files = list(company_dir.glob("*"))
        if not files:
            continue
        
        # Parse dates and get most recent
        dates = set()
        file_list = []
        for f in files:
            if f.name.startswith("."):
                continue
            parts = f.name.split("--")
            if len(parts) >= 2:
                dates.add(parts[0])
            file_list.append({
                "name": f.name,
                "path": str(f),
                "type": get_file_type(f.name),
                "size": f.stat().st_size
            })
        
        # Try to get company/role from job description
        jd_files = list(company_dir.glob("*--job-description.txt"))
        company_name = company_dir.name.replace("-", " ").title()
        role = ""
        url = ""
        
        if jd_files:
            jd_content = jd_files[0].read_text()
            for line in jd_content.split("\n")[:5]:
                if line.startswith("Role:"):
                    role = line.replace("Role:", "").strip()
                elif line.startswith("URL:"):
                    url = line.replace("URL:", "").strip()
        
        applications.append({
            "slug": company_dir.name,
            "company": company_name,
            "role": role,
            "url": url,
            "dates": sorted(dates, reverse=True),
            "files": sorted(file_list, key=lambda x: x["name"], reverse=True),
            "path": str(company_dir)
        })
    
    return applications


def get_file_type(filename: str) -> str:
    """Determine file type from filename"""
    if "resume" in filename:
        return "resume"
    elif "cover-letter" in filename:
        return "cover_letter"
    elif "job-description" in filename:
        return "job_description"
    elif "questions" in filename:
        return "questions"
    return "other"


def scrape_job_posting(url: str) -> Optional[dict]:
    """Scrape job description from a URL. Returns dict with 'content', 'company', 'role'."""
    if not HAS_SCRAPING:
        return None
    
    headers = {
        'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
    }
    
    try:
        response = requests.get(url, headers=headers, timeout=15)
        response.raise_for_status()
    except requests.RequestException as e:
        return {"error": str(e)}
    
    soup = BeautifulSoup(response.text, 'html.parser')
    
    # Remove script, style, nav, footer elements
    for tag in soup(['script', 'style', 'nav', 'footer', 'header', 'aside']):
        tag.decompose()
    
    company = None
    role = None
    
    # Look for title tag
    title_tag = soup.find('title')
    if title_tag:
        title_text = title_tag.get_text()
        if ' at ' in title_text:
            parts = title_text.split(' at ', 1)
            role = parts[0].strip()
            company = parts[1].split('|')[0].split('-')[0].strip()
        elif ' - ' in title_text:
            parts = title_text.split(' - ')
            role = parts[0].strip()
            if len(parts) > 1:
                company = parts[1].strip()
        elif ' | ' in title_text:
            parts = title_text.split(' | ')
            role = parts[0].strip()
            if len(parts) > 1:
                company = parts[1].strip()
    
    # Look for structured data (JSON-LD)
    json_ld = soup.find('script', type='application/ld+json')
    if json_ld:
        try:
            data = json.loads(json_ld.string)
            if isinstance(data, list):
                data = data[0]
            if data.get('@type') == 'JobPosting':
                role = role or data.get('title')
                if data.get('hiringOrganization'):
                    company = company or data['hiringOrganization'].get('name')
        except (json.JSONDecodeError, KeyError, TypeError):
            pass
    
    # Extract main content
    content_selectors = [
        'article',
        '[class*="job-description"]',
        '[class*="jobDescription"]',
        '[class*="job-content"]',
        '[class*="posting-"]',
        '[data-automation*="job"]',
        'main',
        '.content',
        '#content',
    ]
    
    main_content = None
    for selector in content_selectors:
        try:
            main_content = soup.select_one(selector)
            if main_content and len(main_content.get_text(strip=True)) > 200:
                break
        except:
            continue
    
    if not main_content:
        main_content = soup.body if soup.body else soup
    
    text = main_content.get_text(separator='\n', strip=True)
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    content = '\n'.join(lines)
    
    if len(content) > 15000:
        content = content[:15000] + "\n\n[Content truncated...]"
    
    if len(content) < 100:
        return {"error": "Content too short - page may require JavaScript"}
    
    return {
        'content': content,
        'company': company or '',
        'role': role or '',
        'url': url
    }


def tailor_resume_with_ai(master_resume: str, job_description: str, company: str, role: str) -> str:
    """Use Claude to tailor the resume"""
    if not Anthropic:
        raise ImportError("anthropic package not installed")
    
    client = Anthropic(api_key=get_api_key())
    
    prompt = f"""You are tailoring a resume for a specific job application.

CONTENT RULES:
1. KEYWORD ALIGNMENT: Use terms from the JD where the candidate has genuine experience
2. REORDER BULLETS: Put the most relevant accomplishments first  
3. PRESERVE METRICS: Always keep $250K savings and 90 min to 3 min response time
4. NO FABRICATION: Only use keywords where there's real experience
5. ACTIVE VOICE: No "exposure to" or "familiar with"

FORMATTING RULES (CRITICAL - follow exactly):
1. Section headers MUST be ALL CAPS on their own line (e.g., SUMMARY, PROFESSIONAL EXPERIENCE, TECHNICAL SKILLS, PROJECTS, EDUCATION)
2. Company names: Regular case, on their own line (e.g., "Zapier | Remote")
3. Job titles: Regular case, on their own line (e.g., "Incident Manager - Engineering")
4. Dates: On their own line in italics-style, e.g., "Feb 2025 - Present"
5. Bullet points: Start with a hyphen and space "- "
6. NO markdown symbols (no #, **, *, etc.)
7. NO bold markers - the formatting will be applied automatically

MASTER RESUME:
{master_resume}

JOB DESCRIPTION:
{job_description}

COMPANY: {company}
ROLE: {role}

Output the tailored resume in this EXACT structure:

[Name]
[Contact info on one line: email | phone | location | linkedin]

SUMMARY
[2-3 sentence summary]

PROFESSIONAL EXPERIENCE

[Company Name] | [Location]
[Job Title]
[Dates]
- [Bullet point 1]
- [Bullet point 2]
...

TECHNICAL SKILLS
[Category]: [skill1, skill2, skill3]
...

PROJECTS
[Project Name]
- [Description/achievement]
...

EDUCATION
[Degree] - [Institution] ([Years])

Output ONLY the resume content, no explanations."""

    response = client.messages.create(
        model="claude-sonnet-4-20250514",
        max_tokens=2500,
        messages=[{"role": "user", "content": prompt}]
    )
    
    return response.content[0].text.strip()


def generate_cover_letter_with_ai(master_resume: str, job_description: str, company: str, role: str) -> str:
    """Use Claude to generate cover letter"""
    if not Anthropic:
        raise ImportError("anthropic package not installed")
    
    client = Anthropic(api_key=get_api_key())
    
    prompt = f"""You are writing a cover letter for a job application.

CONTEXT:
- Candidate: Ryan Laxson  
- Company: {company}
- Role: {role}

CANDIDATE'S RESUME:
{master_resume}

JOB DESCRIPTION:
{job_description}

COVER LETTER REQUIREMENTS:
1. Start with "Hi [appropriate name from JD] and team," (research who's hiring)
2. Opening: State the role, show genuine interest (2-3 sentences)
3. Body 1: Connect 2-3 specific accomplishments to their needs, include a metric
4. Body 2: Why this specific company (reference something specific, not generic)
5. Closing: Brief call to action, thank them
6. Keep under 350 words total
7. Tone: Professional but warm, authentic not robotic
8. End with "Sincerely," and the name

Output ONLY the cover letter content, no explanations."""

    response = client.messages.create(
        model="claude-sonnet-4-20250514",
        max_tokens=1000,
        messages=[{"role": "user", "content": prompt}]
    )
    
    return response.content[0].text.strip()


def generate_application(company: str, role: str, job_description: str, url: str = None) -> dict:
    """Generate tailored resume and cover letter, save to company folder"""
    master_resume = read_master_resume()
    
    # Generate content
    resume = tailor_resume_with_ai(master_resume, job_description, company, role)
    cover_letter = generate_cover_letter_with_ai(master_resume, job_description, company, role)
    
    # Create company folder
    date_str = datetime.now().strftime("%Y%m%d")
    company_slug = get_company_slug(company)
    company_dir = get_companies_dir() / company_slug
    company_dir.mkdir(parents=True, exist_ok=True)
    
    # Generate filenames
    resume_docx = company_dir / f"{date_str}--resume.docx"
    cover_docx = company_dir / f"{date_str}--cover-letter.docx"
    resume_txt = company_dir / f"{date_str}--resume.txt"
    cover_txt = company_dir / f"{date_str}--cover-letter.txt"
    jd_file = company_dir / f"{date_str}--job-description.txt"
    
    # Save job description
    jd_content = f"JOB DESCRIPTION - {company.upper()}\n"
    jd_content += f"Role: {role}\n"
    jd_content += f"Date: {datetime.now().strftime('%Y-%m-%d')}\n"
    if url:
        jd_content += f"URL: {url}\n"
    jd_content += "=" * 60 + "\n\n"
    jd_content += job_description
    
    with open(jd_file, "w") as f:
        f.write(jd_content)
    
    # Save plain text versions
    with open(resume_txt, "w") as f:
        f.write(resume)
    
    with open(cover_txt, "w") as f:
        f.write(cover_letter)
    
    # Create .docx files
    if Document:
        create_docx(resume, str(resume_docx))
        create_docx(cover_letter, str(cover_docx))
    
    return {
        "company": company,
        "role": role,
        "slug": company_slug,
        "files": {
            "resume_docx": str(resume_docx),
            "resume_txt": str(resume_txt),
            "cover_docx": str(cover_docx),
            "cover_txt": str(cover_txt),
            "job_description": str(jd_file)
        },
        "content": {
            "resume": resume,
            "cover_letter": cover_letter
        }
    }


def strip_markdown_headings(line: str) -> str:
    """Remove markdown heading markers"""
    return re.sub(r'^#{1,6}\s*', '', line)


def strip_markdown_links(text: str) -> str:
    """Convert [text](url) to just text"""
    return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)


def add_formatted_text(paragraph, text: str, font_size: int = 10, base_bold: bool = False, base_italic: bool = False, color = None):
    """Add text to paragraph with markdown bold/italic converted to proper formatting"""
    text = strip_markdown_links(text)
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|[^*]+)'
    segments = re.findall(pattern, text)
    
    for segment in segments:
        if not segment:
            continue
            
        is_bold = base_bold
        is_italic = base_italic
        clean_text = segment
        
        if segment.startswith('**') and segment.endswith('**'):
            is_bold = True
            clean_text = segment[2:-2]
        elif segment.startswith('*') and segment.endswith('*') and len(segment) > 2:
            is_italic = True
            clean_text = segment[1:-1]
        
        if clean_text:
            run = paragraph.add_run(clean_text)
            run.font.size = Pt(font_size)
            run.bold = is_bold
            run.italic = is_italic
            if color:
                run.font.color.rgb = color


SECTION_HEADERS = {
    "SUMMARY", "PROFESSIONAL SUMMARY", "PROFILE",
    "PROFESSIONAL EXPERIENCE", "EXPERIENCE", "WORK EXPERIENCE", "EMPLOYMENT",
    "TECHNICAL SKILLS", "SKILLS", "CORE COMPETENCIES", "TECHNOLOGIES",
    "PROJECTS", "SELECTED PROJECTS", "KEY PROJECTS",
    "EDUCATION", "ACADEMIC BACKGROUND",
    "AWARDS", "HONORS", "ACHIEVEMENTS", "CERTIFICATIONS",
}


def is_section_header(line: str) -> bool:
    """Check if a line is a resume section header"""
    clean = line.strip().upper()
    if clean in SECTION_HEADERS:
        return True
    if line.isupper() and len(line) > 3 and len(line) < 40 and not line.startswith("-"):
        return True
    return False


def is_company_line(line: str) -> bool:
    """Check if a line is a company name line (Company | Location format)"""
    if line.startswith("-"):
        return False
    if "|" in line and any(loc in line for loc in ["Remote", "CA", "NY", "TX", "WA", "OR", "CO", "Fully"]):
        return True
    return False


def is_job_title(line: str) -> bool:
    """Check if a line is a job title"""
    if line.startswith("-"):
        return False
    title_keywords = ["Manager", "Engineer", "Specialist", "Lead", "Director", "Analyst", 
                      "Developer", "Advisor", "Coordinator", "Owner", "Architect"]
    return any(title in line for title in title_keywords) and len(line) < 60


def is_date_line(line: str) -> bool:
    """Check if a line is a date range"""
    if line.startswith("-"):
        return False
    date_indicators = ["Present", "2019", "2020", "2021", "2022", "2023", "2024", "2025", "2026",
                       "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    has_date = any(d in line for d in date_indicators)
    has_separator = "–" in line or "-" in line or "to" in line.lower()
    return has_date and has_separator and len(line) < 50


def create_docx(content: str, filename: str):
    """Create a nicely formatted .docx file"""
    if not Document:
        return
    
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    lines = content.split("\n")
    i = 0
    first_line = True
    
    while i < len(lines):
        line = lines[i].strip()
        
        if not line:
            i += 1
            continue
        
        if line.startswith("===") or line.startswith("---") or line.startswith("___"):
            i += 1
            continue
        
        clean_line = strip_markdown_headings(line)
        p = doc.add_paragraph()
        
        # Name (first non-empty line)
        if first_line and not clean_line.startswith("-"):
            add_formatted_text(p, clean_line, font_size=18, base_bold=True, color=RGBColor(0, 51, 102))
            p.paragraph_format.space_after = Pt(2)
            first_line = False
        
        # Contact info (contains | or @ near top of document)
        elif ("|" in clean_line or "@" in clean_line) and i < 5:
            add_formatted_text(p, clean_line, font_size=10, color=RGBColor(80, 80, 80))
            p.paragraph_format.space_after = Pt(12)
        
        # Section headers (PROFESSIONAL EXPERIENCE, SKILLS, etc.)
        elif is_section_header(clean_line):
            header_text = clean_line.upper()
            add_formatted_text(p, header_text, font_size=11, base_bold=True, color=RGBColor(0, 51, 102))
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        
        # Markdown headers (## Header)
        elif line.startswith("## "):
            header_text = clean_line.upper()
            add_formatted_text(p, header_text, font_size=11, base_bold=True, color=RGBColor(0, 51, 102))
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        
        # Sub-headers (### for company names in markdown)
        elif line.startswith("### "):
            add_formatted_text(p, clean_line, font_size=11, base_bold=True)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
        
        # Company lines (Zapier | Remote)
        elif is_company_line(clean_line):
            add_formatted_text(p, clean_line, font_size=11, base_bold=True)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
        
        # Job titles (Incident Manager, Senior Engineer, etc.)
        elif is_job_title(clean_line):
            add_formatted_text(p, clean_line, font_size=10, base_bold=True)
            p.paragraph_format.space_after = Pt(1)
        
        # Date lines (Feb 2025 - Present)
        elif is_date_line(clean_line):
            date_text = clean_line.strip("*")
            run = p.add_run(date_text)
            run.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(100, 100, 100)
            p.paragraph_format.space_after = Pt(6)
        
        # Bullet points
        elif clean_line.startswith("-") or clean_line.startswith("•"):
            bullet_text = clean_line.lstrip("-•").strip()
            p.add_run("• ")
            add_formatted_text(p, bullet_text, font_size=10)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(4)
        
        # Labeled lines (Skills: Python, JavaScript)
        elif ":" in clean_line and not clean_line.startswith("-") and "|" not in clean_line and len(clean_line.split(":")[0]) < 30:
            parts = clean_line.split(":", 1)
            if len(parts) == 2 and parts[1].strip():
                label = parts[0].strip("*")
                run1 = p.add_run(label + ":")
                run1.bold = True
                run1.font.size = Pt(10)
                add_formatted_text(p, " " + parts[1].strip(), font_size=10)
                p.paragraph_format.space_after = Pt(3)
            else:
                add_formatted_text(p, clean_line, font_size=10)
                p.paragraph_format.space_after = Pt(4)
        
        # Project names
        elif any(kw in clean_line for kw in ["Bot", "Framework", "Platform", "System", "Tool"]) and len(clean_line) < 60 and not clean_line.startswith("-"):
            add_formatted_text(p, clean_line, font_size=10, base_bold=True)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        
        elif len(clean_line) > 100:
            add_formatted_text(p, clean_line, font_size=10)
            p.paragraph_format.space_after = Pt(8)
        
        else:
            add_formatted_text(p, clean_line, font_size=10)
            p.paragraph_format.space_after = Pt(4)
        
        i += 1
    
    doc.save(filename)
